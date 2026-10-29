#!/usr/bin/env python3
"""
Helper script to extract text and tables from docx and pdf files in lab docs.
Usage:
    python3 extract_doc.py <path_to_file> [output_file]
"""
import sys
import os

def extract(path, output_path=None):
    if not os.path.exists(path):
        print(f"Error: File {path} not found.")
        sys.exit(1)
        
    ext = os.path.splitext(path)[1].lower()
    text_content = []

    if ext == ".docx":
        try:
            import docx
            doc = docx.Document(path)
            for p in doc.paragraphs:
                if p.text.strip():
                    text_content.append(p.text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells)
                    if row_text.strip():
                        text_content.append(row_text)
        except Exception as e:
            print(f"Error extracting docx: {e}")
            sys.exit(1)

    elif ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(path)
            for i, page in enumerate(reader.pages):
                text_content.append(f"--- PAGE {i+1} ---")
                text_content.append(page.extract_text() or "")
        except Exception as e:
            print(f"Error extracting pdf: {e}")
            sys.exit(1)
    else:
        print(f"Unsupported format: {ext}")
        sys.exit(1)

    full_text = "\n".join(text_content)
    try:
        if output_path:
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(full_text)
            print(f"Successfully extracted text to {output_path}")
        else:
            print(full_text)
    except BrokenPipeError:
        devnull = os.open(os.devnull, os.O_WRONLY)
        os.dup2(devnull, sys.stdout.fileno())
        sys.exit(0)

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python3 extract_doc.py <path_to_file> [output_file]")
        sys.exit(1)
    out = sys.argv[2] if len(sys.argv) > 2 else None
    extract(sys.argv[1], out)
